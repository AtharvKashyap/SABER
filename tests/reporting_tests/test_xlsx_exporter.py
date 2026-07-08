"""Tests for XLSX report exporter."""

from __future__ import annotations

from openpyxl import load_workbook

from saber.reporting.json_exporter import (
    ReportDocument,
    ReportFinding,
    ReportObservation,
    ReportSeverity,
)
from saber.reporting.xlsx_exporter import XlsxExporter


def make_document() -> ReportDocument:
    """Create report document."""

    return ReportDocument(
        report_id="report_xlsx_test",
        mission_name="XLSX Mission",
        target="example.com",
        generated_at=__import__("datetime").datetime(2026, 1, 1, 12, 0, 0, tzinfo=__import__("datetime").UTC),
        executive_summary="XLSX summary.",
        observations=[
            ReportObservation(
                kind="service",
                summary="Port 443 open.",
                source_tool="nmap",
                data={"host": "10.0.0.5", "port": 443},
                metadata={"format": "xml"},
            ),
            ReportObservation(
                kind="web_technology",
                summary="Uses nginx.",
                source_tool="whatweb",
                data={"server": "nginx"},
                metadata={"format": "json"},
            ),
        ],
        findings=[
            ReportFinding(
                title="Low Finding",
                severity=ReportSeverity.LOW,
                description="Low description.",
                source_tool="nuclei",
                evidence={"template_id": "low-template"},
                references=[],
                metadata={},
            ),
            ReportFinding(
                title="Critical Finding",
                severity=ReportSeverity.CRITICAL,
                description="Critical description.",
                source_tool="nuclei",
                evidence={"template_id": "critical-template", "nested": {"a": 1}},
                references=["https://example.com/critical"],
                metadata={"tag": "critical"},
            ),
        ],
        metadata={"owner": "unit-test", "scope": {"approved": True}},
    )


class TestXlsxExporter:
    """Validate XLSX exporter."""

    def test_export_creates_workbook(self, tmp_path) -> None:
        """Exporter should create XLSX workbook."""

        output_path = tmp_path / "report.xlsx"

        returned = XlsxExporter().export(make_document(), output_path)

        assert returned == output_path
        assert output_path.exists()
        assert output_path.stat().st_size > 0

    def test_workbook_sheets(self, tmp_path) -> None:
        """Workbook should include expected sheets."""

        output_path = tmp_path / "report.xlsx"
        XlsxExporter().export(make_document(), output_path)

        workbook = load_workbook(output_path)

        assert workbook.sheetnames == ["Summary", "Findings", "Observations", "Evidence", "Metadata"]

    def test_summary_sheet_content(self, tmp_path) -> None:
        """Summary sheet should contain mission metadata and counts."""

        output_path = tmp_path / "report.xlsx"
        XlsxExporter().export(make_document(), output_path)

        sheet = load_workbook(output_path)["Summary"]
        rows = {row[0].value: row[1].value for row in sheet.iter_rows(min_row=2, max_col=2)}

        assert rows["Report ID"] == "report_xlsx_test"
        assert rows["Mission Name"] == "XLSX Mission"
        assert rows["Target"] == "example.com"
        assert rows["Executive Summary"] == "XLSX summary."
        assert rows["Total Findings"] == 2
        assert rows["Total Observations"] == 2
        assert rows["Highest Severity"] == "critical"
        assert rows["Critical"] == 1
        assert rows["Low"] == 1

    def test_findings_sheet_content_and_sorting(self, tmp_path) -> None:
        """Findings should be sorted by severity."""

        output_path = tmp_path / "report.xlsx"
        XlsxExporter().export(make_document(), output_path)

        sheet = load_workbook(output_path)["Findings"]

        headers = [cell.value for cell in sheet[1]]
        assert headers == ["Severity", "Title", "Description", "Source Tool", "References", "Evidence JSON", "Metadata JSON"]

        assert sheet["A2"].value == "critical"
        assert sheet["B2"].value == "Critical Finding"
        assert sheet["D2"].value == "nuclei"
        assert sheet["E2"].value == "https://example.com/critical"
        assert '"template_id": "critical-template"' in sheet["F2"].value
        assert '"tag": "critical"' in sheet["G2"].value

        assert sheet["A3"].value == "low"
        assert sheet["B3"].value == "Low Finding"

    def test_observations_sheet_content(self, tmp_path) -> None:
        """Observations sheet should contain normalized observations."""

        output_path = tmp_path / "report.xlsx"
        XlsxExporter().export(make_document(), output_path)

        sheet = load_workbook(output_path)["Observations"]

        headers = [cell.value for cell in sheet[1]]
        assert headers == ["Kind", "Summary", "Source Tool", "Data JSON", "Metadata JSON"]

        assert sheet["A2"].value == "service"
        assert sheet["B2"].value == "Port 443 open."
        assert sheet["C2"].value == "nmap"
        assert '"port": 443' in sheet["D2"].value

        assert sheet["A3"].value == "web_technology"
        assert sheet["C3"].value == "whatweb"
        assert '"server": "nginx"' in sheet["D3"].value

    def test_evidence_sheet_content(self, tmp_path) -> None:
        """Evidence sheet should flatten finding evidence."""

        output_path = tmp_path / "report.xlsx"
        XlsxExporter().export(make_document(), output_path)

        sheet = load_workbook(output_path)["Evidence"]

        headers = [cell.value for cell in sheet[1]]
        assert headers == ["Finding Title", "Severity", "Evidence Key", "Evidence Value"]

        values = [
            (row[0].value, row[1].value, row[2].value, row[3].value)
            for row in sheet.iter_rows(min_row=2, max_col=4)
        ]

        assert ("Critical Finding", "critical", "template_id", "critical-template") in values
        assert any(row[0] == "Critical Finding" and row[2] == "nested" and '"a": 1' in row[3] for row in values)
        assert ("Low Finding", "low", "template_id", "low-template") in values

    def test_metadata_sheet_content(self, tmp_path) -> None:
        """Metadata sheet should contain document metadata."""

        output_path = tmp_path / "report.xlsx"
        XlsxExporter().export(make_document(), output_path)

        sheet = load_workbook(output_path)["Metadata"]
        rows = {row[0].value: row[1].value for row in sheet.iter_rows(min_row=2, max_col=2)}

        assert rows["owner"] == "unit-test"
        assert '"approved": true' in rows["scope"]

    def test_header_style_and_freeze_panes(self, tmp_path) -> None:
        """Sheets should have frozen header rows and bold headers."""

        output_path = tmp_path / "report.xlsx"
        XlsxExporter().export(make_document(), output_path)

        workbook = load_workbook(output_path)

        for sheet in workbook.worksheets:
            assert sheet.freeze_panes == "A2"
            assert sheet["A1"].font.bold is True
            assert sheet.column_dimensions["A"].width >= 12

    def test_export_empty_document(self, tmp_path) -> None:
        """Exporter should handle empty findings and observations."""

        document = ReportDocument(
            report_id="empty",
            mission_name="Empty Mission",
            target="example.com",
            generated_at=__import__("datetime").datetime(2026, 1, 1, 12, 0, 0, tzinfo=__import__("datetime").UTC),
            executive_summary="Empty summary.",
            observations=[],
            findings=[],
            metadata={},
        )
        output_path = tmp_path / "empty.xlsx"

        XlsxExporter().export(document, output_path)
        workbook = load_workbook(output_path)

        assert workbook["Summary"]["B6"].value == 0
        assert workbook["Findings"].max_row == 1
        assert workbook["Observations"].max_row == 1
        assert workbook["Evidence"].max_row == 1
        assert workbook["Metadata"].max_row == 1
