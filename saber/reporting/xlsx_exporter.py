"""XLSX report exporter for SABER."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from saber.reporting.json_exporter import ReportDocument, ReportSeverity, SEVERITY_ORDER


class XlsxExporter:
    """Export SABER report documents to XLSX workbooks."""

    def export(self, document: ReportDocument, output_path: str | Path) -> Path:
        """Write report document to XLSX file."""

        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"

        self._write_summary(summary, document)
        self._write_findings(workbook.create_sheet("Findings"), document)
        self._write_observations(workbook.create_sheet("Observations"), document)
        self._write_evidence(workbook.create_sheet("Evidence"), document)
        self._write_metadata(workbook.create_sheet("Metadata"), document)

        for sheet in workbook.worksheets:
            self._style_sheet(sheet)

        workbook.save(path)
        return path

    def _write_summary(self, sheet: Worksheet, document: ReportDocument) -> None:
        """Write summary sheet."""

        counts = document.severity_counts()
        rows = [
            ("Report ID", document.report_id),
            ("Mission Name", document.mission_name),
            ("Target", document.target),
            ("Generated At", document.generated_at.isoformat()),
            ("Executive Summary", document.executive_summary),
            ("Total Findings", len(document.findings)),
            ("Total Observations", len(document.observations)),
            ("Highest Severity", document.highest_severity().value),
            ("Critical", counts[ReportSeverity.CRITICAL.value]),
            ("High", counts[ReportSeverity.HIGH.value]),
            ("Medium", counts[ReportSeverity.MEDIUM.value]),
            ("Low", counts[ReportSeverity.LOW.value]),
            ("Info", counts[ReportSeverity.INFO.value]),
            ("Unknown", counts[ReportSeverity.UNKNOWN.value]),
        ]

        sheet.append(["Field", "Value"])
        for row in rows:
            sheet.append(list(row))

    def _write_findings(self, sheet: Worksheet, document: ReportDocument) -> None:
        """Write findings sheet."""

        sheet.append(["Severity", "Title", "Description", "Source Tool", "References", "Evidence JSON", "Metadata JSON"])

        for finding in sorted(document.findings, key=lambda item: SEVERITY_ORDER[item.severity]):
            sheet.append(
                [
                    finding.severity.value,
                    finding.title,
                    finding.description,
                    finding.source_tool or "",
                    "\n".join(finding.references),
                    self._json(finding.evidence),
                    self._json(finding.metadata),
                ]
            )

    def _write_observations(self, sheet: Worksheet, document: ReportDocument) -> None:
        """Write observations sheet."""

        sheet.append(["Kind", "Summary", "Source Tool", "Data JSON", "Metadata JSON"])

        for observation in document.observations:
            sheet.append(
                [
                    observation.kind,
                    observation.summary,
                    observation.source_tool or "",
                    self._json(observation.data),
                    self._json(observation.metadata),
                ]
            )

    def _write_evidence(self, sheet: Worksheet, document: ReportDocument) -> None:
        """Write flattened finding evidence sheet."""

        sheet.append(["Finding Title", "Severity", "Evidence Key", "Evidence Value"])

        for finding in sorted(document.findings, key=lambda item: SEVERITY_ORDER[item.severity]):
            if not finding.evidence:
                sheet.append([finding.title, finding.severity.value, "", ""])
                continue

            for key, value in finding.evidence.items():
                sheet.append([finding.title, finding.severity.value, key, self._stringify(value)])

    def _write_metadata(self, sheet: Worksheet, document: ReportDocument) -> None:
        """Write metadata sheet."""

        sheet.append(["Key", "Value"])

        if not document.metadata:
            return

        for key, value in document.metadata.items():
            sheet.append([key, self._stringify(value)])

    def _style_sheet(self, sheet: Worksheet) -> None:
        """Apply basic workbook styling."""

        sheet.freeze_panes = "A2"

        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        self._autosize_columns(sheet)

    @staticmethod
    def _autosize_columns(sheet: Worksheet) -> None:
        """Best-effort autosize columns."""

        for column_cells in sheet.columns:
            column_letter = column_cells[0].column_letter
            max_length = 0

            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, min(len(value), 80))

            sheet.column_dimensions[column_letter].width = max(max_length + 2, 12)

    @staticmethod
    def _json(value: Any) -> str:
        """Serialize value as pretty JSON."""

        return json.dumps(value, indent=2, sort_keys=True, default=str)

    @staticmethod
    def _stringify(value: Any) -> str:
        """Stringify scalar or complex value."""

        if isinstance(value, (dict, list, tuple)):
            return json.dumps(value, indent=2, sort_keys=True, default=str)
        return "" if value is None else str(value)
